# Copyright 2026 Amazon.com, Inc. or its affiliates. All Rights Reserved.
# SPDX-License-Identifier: MIT-0

"""PPTX native parsing, slide classification, and the sidecar client.

This module was at 0% coverage despite owning the whole PowerPoint ingest path.
Everything here runs without LibreOffice, without AWS and without network: the
parsing tests build real .pptx files in memory with python-pptx, and the sidecar
tests stub httpx.post.

The classification thresholds only choose a prompt template (both text and image
go to the model regardless), so the tests pin the boundaries rather than
asserting any particular slide "should" be one type.
"""

from __future__ import annotations

import io

import pytest

from app.models.document import SlideType
from app.services import pptx_converter as pc


# ---------------------------------------------------------------------------
# Deck builders — real .pptx bytes, no fixtures on disk
# ---------------------------------------------------------------------------

def _blank_layout(pres):
    """The 'Blank' layout, so no placeholder shapes are added for us."""
    return pres.slide_layouts[6]


def _deck(build) -> bytes:
    """Build a presentation via `build(pres)` and return its bytes."""
    from pptx import Presentation

    pres = Presentation()
    build(pres)
    buf = io.BytesIO()
    pres.save(buf)
    return buf.getvalue()


def _add_textbox(slide, text: str, top_in: float = 1.0):
    from pptx.util import Inches

    box = slide.shapes.add_textbox(Inches(1), Inches(top_in), Inches(6), Inches(1))
    box.text_frame.text = text
    return box


def _png_bytes() -> bytes:
    """A 1x1 PNG — enough for python-pptx to embed a real picture shape."""
    import base64

    return base64.b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8AAAwAB/wFxrx4AAAAASUVORK5CYII="
    )


class TestParseNative:
    def test_reads_text_from_every_slide(self):
        def build(pres):
            for text in ("First slide", "Second slide"):
                slide = pres.slides.add_slide(_blank_layout(pres))
                _add_textbox(slide, text)

        meta = pc.parse_pptx_native(_deck(build))
        assert meta.slide_count == 2
        assert [s.slide_index for s in meta.slides] == [1, 2]
        assert "First slide" in meta.slides[0].native_text
        assert "Second slide" in meta.slides[1].native_text

    def test_slide_index_is_one_based(self):
        """The S3 slide-image keys are built from this, so an off-by-one here
        would point every slide at the wrong PNG."""
        def build(pres):
            _add_textbox(pres.slides.add_slide(_blank_layout(pres)), "only")

        assert pc.parse_pptx_native(_deck(build)).slides[0].slide_index == 1

    def test_table_text_is_folded_into_native_text(self):
        """The model reads native_text; a table whose wording only lived in a
        separate field would be invisible to rule extraction."""
        def build(pres):
            from pptx.util import Inches

            slide = pres.slides.add_slide(_blank_layout(pres))
            shape = slide.shapes.add_table(2, 2, Inches(1), Inches(1), Inches(6), Inches(2))
            cells = shape.table
            cells.cell(0, 0).text = "Region"
            cells.cell(0, 1).text = "AZ count"
            cells.cell(1, 0).text = "eu-west-1"
            cells.cell(1, 1).text = "3"

        slide = pc.parse_pptx_native(_deck(build)).slides[0]
        assert slide.table_count == 1
        assert slide.tables[0][0] == ["Region", "AZ count"]
        assert slide.tables[0][1] == ["eu-west-1", "3"]
        assert "eu-west-1" in slide.native_text
        assert "AZ count" in slide.native_text

    def test_counts_pictures_separately_from_text_shapes(self):
        def build(pres):
            from pptx.util import Inches

            slide = pres.slides.add_slide(_blank_layout(pres))
            _add_textbox(slide, "caption")
            slide.shapes.add_picture(
                io.BytesIO(_png_bytes()), Inches(1), Inches(3), Inches(1), Inches(1)
            )

        slide = pc.parse_pptx_native(_deck(build)).slides[0]
        assert slide.image_count == 1
        assert slide.shape_count == 2
        assert "caption" in slide.native_text

    def test_descends_into_grouped_shapes(self):
        """Architecture diagrams are usually groups; not descending would report
        one shape and misclassify the slide as title-only."""
        def build(pres):
            from pptx.util import Inches

            slide = pres.slides.add_slide(_blank_layout(pres))
            a = _add_textbox(slide, "inner one", top_in=1.0)
            b = _add_textbox(slide, "inner two", top_in=2.5)
            slide.shapes.add_group_shape([a, b])

        slide = pc.parse_pptx_native(_deck(build)).slides[0]
        assert "inner one" in slide.native_text
        assert "inner two" in slide.native_text
        # The group itself is not counted, only its leaves.
        assert slide.shape_count == 2

    def test_hidden_slides_are_skipped(self):
        """LibreOffice's PDF export drops hidden slides, so counting them here
        would desync slide_count from the PDF page_count and trip the upload
        route's alignment guard."""
        def build(pres):
            visible = pres.slides.add_slide(_blank_layout(pres))
            _add_textbox(visible, "shown")
            hidden = pres.slides.add_slide(_blank_layout(pres))
            _add_textbox(hidden, "not shown")
            hidden._element.set("show", "0")

        meta = pc.parse_pptx_native(_deck(build))
        assert meta.slide_count == 1
        assert "shown" in meta.slides[0].native_text
        assert "not shown" not in meta.slides[0].native_text

    def test_empty_deck_yields_no_slides(self):
        assert pc.parse_pptx_native(_deck(lambda pres: None)).slide_count == 0

    def test_corrupt_bytes_raise(self):
        with pytest.raises(Exception):
            pc.parse_pptx_native(b"this is not a pptx file")

    def test_every_slide_gets_a_type(self):
        def build(pres):
            _add_textbox(pres.slides.add_slide(_blank_layout(pres)), "x")

        assert pc.parse_pptx_native(_deck(build)).slides[0].slide_type in set(SlideType)


# ---------------------------------------------------------------------------
# Classification thresholds
# ---------------------------------------------------------------------------

def _native(text: str = "", shapes: int = 1, images: int = 0) -> pc.SlideNative:
    return pc.SlideNative(
        slide_index=1, native_text=text, shape_count=shapes, image_count=images
    )


class TestClassifySlide:
    def test_short_text_and_few_shapes_is_title_only(self):
        assert pc.classify_slide(_native("Agenda", shapes=1)) is SlideType.TITLE_ONLY

    def test_title_only_boundary_is_inclusive(self):
        at_limit = _native("x" * pc.TITLE_TEXT_MAX, shapes=pc.TITLE_SHAPE_MAX)
        assert pc.classify_slide(at_limit) is SlideType.TITLE_ONLY

    def test_one_shape_over_the_title_limit_is_not_title_only(self):
        over = _native("x" * pc.TITLE_TEXT_MAX, shapes=pc.TITLE_SHAPE_MAX + 1)
        assert pc.classify_slide(over) is not SlideType.TITLE_ONLY

    def test_long_text_with_no_images_is_text_heavy(self):
        slide = _native("x" * pc.TEXT_HEAVY_TEXT_MIN, shapes=2, images=0)
        assert pc.classify_slide(slide) is SlideType.TEXT_HEAVY

    def test_a_single_image_disqualifies_text_heavy(self):
        """text_heavy drives the verbatim-wording prompt; an image means the
        slide also carries visual content the other template handles."""
        slide = _native("x" * pc.TEXT_HEAVY_TEXT_MIN, shapes=2, images=1)
        assert pc.classify_slide(slide) is not SlideType.TEXT_HEAVY

    def test_many_shapes_with_little_text_is_diagram_heavy(self):
        slide = _native("VPC", shapes=pc.DIAGRAM_HEAVY_SHAPE_MIN)
        assert pc.classify_slide(slide) is SlideType.DIAGRAM_HEAVY

    def test_diagram_heavy_needs_the_shape_minimum(self):
        slide = _native("x" * 150, shapes=pc.DIAGRAM_HEAVY_SHAPE_MIN - 1)
        assert pc.classify_slide(slide) is not SlideType.DIAGRAM_HEAVY

    def test_anything_else_falls_back_to_hybrid(self):
        # Mid-length text, mid shape count, has an image: matches no rule.
        slide = _native("x" * 300, shapes=5, images=1)
        assert pc.classify_slide(slide) is SlideType.HYBRID

    def test_classification_never_returns_none(self):
        for text_len in (0, 79, 80, 199, 200, 399, 400, 5000):
            for shapes in (0, 1, 3, 4, 6, 7, 8, 40):
                for images in (0, 1):
                    got = pc.classify_slide(_native("x" * text_len, shapes, images))
                    assert got in set(SlideType)


# ---------------------------------------------------------------------------
# Sidecar client — status-code -> domain-exception mapping
# ---------------------------------------------------------------------------

class _Resp:
    def __init__(self, status: int, body=None, raise_exc=None):
        self.status_code = status
        self._body = body
        self._raise = raise_exc

    def json(self):
        if isinstance(self._body, Exception):
            raise self._body
        if self._body is None:
            raise ValueError("no body")
        return self._body

    def raise_for_status(self):
        if self._raise:
            raise self._raise
        raise RuntimeError(f"unexpected status {self.status_code}")


@pytest.fixture
def post(monkeypatch):
    """Replace httpx.post and record the request it was given."""
    calls: list[dict] = []
    box: dict = {}

    def _fake_post(url, json=None, timeout=None):
        calls.append({"url": url, "json": json, "timeout": timeout})
        result = box["result"]
        if isinstance(result, Exception):
            raise result
        return result

    monkeypatch.setattr(pc.httpx, "post", _fake_post)

    def _set(result):
        box["result"] = result
        return calls

    return _set


class TestConvertPptxToPdf:
    def test_returns_the_page_count_on_success(self, post):
        post(_Resp(200, {"page_count": 52}))
        assert (
            pc.convert_pptx_to_pdf("s3://src", "s3://dst", base_url="http://side")
            == 52
        )

    def test_posts_the_presigned_urls_not_the_bytes(self, post):
        """Bytes must not cross the API<->sidecar boundary; the sidecar reads
        and writes S3 itself via these URLs."""
        calls = post(_Resp(200, {"page_count": 1}))
        pc.convert_pptx_to_pdf("https://src", "https://dst", base_url="http://side")
        assert calls[0]["url"] == "http://side/convert"
        assert calls[0]["json"]["source_url"] == "https://src"
        assert calls[0]["json"]["target_url"] == "https://dst"
        assert calls[0]["timeout"] == pc.SIDECAR_TIMEOUT_S

    def test_strips_a_trailing_slash_from_the_base_url(self, post):
        calls = post(_Resp(200, {"page_count": 1}))
        pc.convert_pptx_to_pdf("a", "b", base_url="http://side/")
        assert calls[0]["url"] == "http://side/convert"

    def test_reads_the_base_url_from_the_environment(self, post, monkeypatch):
        monkeypatch.setenv("PPTX_CONVERTER_URL", "http://from-env:9000")
        calls = post(_Resp(200, {"page_count": 1}))
        pc.convert_pptx_to_pdf("a", "b")
        assert calls[0]["url"] == "http://from-env:9000/convert"

    def test_falls_back_to_the_default_url(self, post, monkeypatch):
        monkeypatch.delenv("PPTX_CONVERTER_URL", raising=False)
        calls = post(_Resp(200, {"page_count": 1}))
        pc.convert_pptx_to_pdf("a", "b")
        assert calls[0]["url"] == f"{pc.DEFAULT_CONVERTER_URL}/convert"

    def test_passes_the_size_cap_through(self, post):
        calls = post(_Resp(200, {"page_count": 1}))
        pc.convert_pptx_to_pdf("a", "b", max_bytes=123, base_url="http://side")
        assert calls[0]["json"]["max_bytes"] == 123

    def test_a_client_timeout_becomes_conversion_timeout(self, post):
        post(pc.httpx.TimeoutException("read timeout"))
        with pytest.raises(pc.ConversionTimeout):
            pc.convert_pptx_to_pdf("a", "b", base_url="http://side")

    def test_504_becomes_conversion_timeout(self, post):
        post(_Resp(504))
        with pytest.raises(pc.ConversionTimeout):
            pc.convert_pptx_to_pdf("a", "b", base_url="http://side")

    @pytest.mark.parametrize("status", [413, 422])
    def test_rejected_input_becomes_conversion_error(self, post, status):
        post(_Resp(status, {"detail": "password-protected deck"}))
        with pytest.raises(pc.ConversionError) as exc:
            pc.convert_pptx_to_pdf("a", "b", base_url="http://side")
        assert "password-protected" in str(exc.value)

    def test_502_becomes_upstream_error(self, post):
        post(_Resp(502, {"detail": "S3 returned 403"}))
        with pytest.raises(pc.ConversionUpstreamError) as exc:
            pc.convert_pptx_to_pdf("a", "b", base_url="http://side")
        assert "S3" in str(exc.value)

    def test_non_json_2xx_is_an_error_not_a_zero_page_count(self, post):
        """Returning 0 would let the caller believe a 0-page PDF was produced."""
        post(_Resp(200, None))
        with pytest.raises(pc.ConversionError):
            pc.convert_pptx_to_pdf("a", "b", base_url="http://side")

    def test_a_missing_detail_still_raises_with_a_message(self, post):
        post(_Resp(422, {}))
        with pytest.raises(pc.ConversionError) as exc:
            pc.convert_pptx_to_pdf("a", "b", base_url="http://side")
        assert "422" in str(exc.value)

    def test_unexpected_status_bubbles_up(self, post):
        post(_Resp(500, {}, raise_exc=RuntimeError("boom")))
        with pytest.raises(RuntimeError):
            pc.convert_pptx_to_pdf("a", "b", base_url="http://side")


class TestSafeDetail:
    def test_extracts_the_detail_field(self):
        assert pc._safe_detail(_Resp(422, {"detail": "nope"})) == "nope"

    def test_non_json_body_yields_empty_string(self):
        assert pc._safe_detail(_Resp(422, None)) == ""

    def test_non_dict_json_yields_empty_string(self):
        assert pc._safe_detail(_Resp(422, ["a", "b"])) == ""


# ---------------------------------------------------------------------------
# Rasterisation — real PDF via PyMuPDF, S3 replaced with a recorder
# ---------------------------------------------------------------------------

class _RecordingStore:
    def __init__(self):
        self.uploads: list[tuple[str, str, int]] = []

    def upload_document(self, data, key, content_type=None):
        self.uploads.append((key, content_type, len(data)))
        return key


def _pdf_bytes(pages: int) -> bytes:
    import fitz

    doc = fitz.open()
    for _ in range(pages):
        doc.new_page()
    out = doc.tobytes()
    doc.close()
    return out


class TestRasteriseSlidesToS3:
    def test_one_png_per_page_in_order(self):
        store = _RecordingStore()
        keys = pc.rasterise_slides_to_s3(_pdf_bytes(3), "doc-1", store)
        assert keys == [
            "documents/doc-1/slides/1.png",
            "documents/doc-1/slides/2.png",
            "documents/doc-1/slides/3.png",
        ]
        assert [u[0] for u in store.uploads] == keys

    def test_slide_numbers_are_one_based(self):
        """These keys are matched against SlideStatus.image_s3_key, which uses
        1-based slide_index."""
        store = _RecordingStore()
        keys = pc.rasterise_slides_to_s3(_pdf_bytes(1), "d", store)
        assert keys[0].endswith("/1.png")

    def test_uploads_are_tagged_as_png(self):
        store = _RecordingStore()
        pc.rasterise_slides_to_s3(_pdf_bytes(1), "d", store)
        assert store.uploads[0][1] == "image/png"
        assert store.uploads[0][2] > 0

    def test_higher_dpi_produces_a_larger_image(self):
        low, high = _RecordingStore(), _RecordingStore()
        pc.rasterise_slides_to_s3(_pdf_bytes(1), "d", low, dpi=72)
        pc.rasterise_slides_to_s3(_pdf_bytes(1), "d", high, dpi=288)
        assert high.uploads[0][2] > low.uploads[0][2]

    def test_unreadable_pdf_bytes_raise_before_any_upload(self):
        """A zero-page PDF is not constructible (PyMuPDF refuses to save one),
        so the empty case in practice is a corrupt stream. It must fail loudly
        rather than report success with no slides."""
        store = _RecordingStore()
        with pytest.raises(Exception):
            pc.rasterise_slides_to_s3(b"not a pdf", "d", store)
        assert store.uploads == []

    def test_an_upload_failure_propagates(self):
        """A silently-swallowed failure would leave slides with no image while
        the document looked fully processed."""
        class _Broken:
            def upload_document(self, *a, **kw):
                raise RuntimeError("S3 down")

        with pytest.raises(RuntimeError):
            pc.rasterise_slides_to_s3(_pdf_bytes(2), "d", _Broken())
