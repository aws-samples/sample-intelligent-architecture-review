# Copyright 2026 Amazon.com, Inc. or its affiliates. All Rights Reserved.
# SPDX-License-Identifier: MIT-0

"""Ingestion: normalising PDF pages and PPTX slides into unified Units.

This is the single front door for extraction, and it was at 9% coverage. The
important behaviour is that everything about the multimodal path degrades rather
than fails: if slide PNGs are missing, if rasterisation blows up, if a page
upload fails, extraction must continue text-only instead of losing the document.

The tests build real .pptx and .pdf bytes in memory (python-pptx, PyMuPDF) and
substitute a recording document store, so no AWS and no sidecar is involved.
"""

from __future__ import annotations

import io

import pytest

from app.agents.extraction import ingestion as ing
from app.models.document import Document, DocumentStatus, DocumentType, SourceFormat


class _DocStore:
    """Recording document store. `present` lists keys that already exist."""

    def __init__(self, present=None, upload_error=None):
        self.present = set(present or [])
        self.uploads: list[tuple[str, int]] = []
        self.upload_error = upload_error

    def download_document(self, key):
        if key not in self.present:
            raise RuntimeError(f"no such key: {key}")
        return b"existing-png"

    def upload_document(self, data, key, content_type=None):
        if self.upload_error:
            raise self.upload_error
        self.uploads.append((key, len(data)))
        return key


def _pptx_bytes(slides=2, with_table=False) -> bytes:
    from pptx import Presentation
    from pptx.util import Inches

    pres = Presentation()
    blank = pres.slide_layouts[6]
    for i in range(1, slides + 1):
        slide = pres.slides.add_slide(blank)
        box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1))
        box.text_frame.text = f"Slide {i} content about availability"
        if with_table and i == 1:
            shape = slide.shapes.add_table(2, 2, Inches(1), Inches(3), Inches(6), Inches(2))
            t = shape.table
            t.cell(0, 0).text = "Primary"
            t.cell(0, 1).text = "DR"
            t.cell(1, 0).text = "VA"
            t.cell(1, 1).text = "TX"
    buf = io.BytesIO()
    pres.save(buf)
    return buf.getvalue()


def _pdf_bytes(pages=2, text="Availability requirements") -> bytes:
    import fitz

    doc = fitz.open()
    for i in range(pages):
        page = doc.new_page()
        page.insert_text((72, 72), f"{text} page {i + 1}")
    out = doc.tobytes()
    doc.close()
    return out


def _doc(source_format=SourceFormat.PDF) -> Document:
    return Document(
        document_id="doc-1",
        filename="f",
        s3_key="k",
        document_type=DocumentType.STANDARDS,
        status=DocumentStatus.PARSED,
        hosting_scope="both",
        source_format=source_format,
    )


class TestIngestPptx:
    def test_one_unit_per_slide(self):
        ds = _DocStore(present=["documents/doc-1/slides/1.png"])
        units = ing.ingest_pptx(_pptx_bytes(3), "doc-1", ds)
        assert len(units) == 3
        assert [u.index for u in units] == [1, 2, 3]
        assert all(u.source_format == "pptx" for u in units)

    def test_native_text_is_carried(self):
        ds = _DocStore(present=["documents/doc-1/slides/1.png"])
        units = ing.ingest_pptx(_pptx_bytes(1), "doc-1", ds)
        assert "availability" in units[0].text

    def test_each_unit_points_at_its_slide_image(self):
        """image_ref is what makes extraction multimodal; losing it silently
        reverts the deck to text-only."""
        ds = _DocStore(present=["documents/doc-1/slides/1.png"])
        units = ing.ingest_pptx(_pptx_bytes(2), "doc-1", ds)
        assert units[0].image_ref == "documents/doc-1/slides/1.png"
        assert units[1].image_ref == "documents/doc-1/slides/2.png"

    def test_structured_tables_are_preserved_as_grids(self):
        ds = _DocStore(present=["documents/doc-1/slides/1.png"])
        units = ing.ingest_pptx(_pptx_bytes(1, with_table=True), "doc-1", ds)
        assert units[0].tables
        assert units[0].tables[0].rows[0] == ["Primary", "DR"]

    def test_classification_is_carried_from_the_parser(self):
        ds = _DocStore(present=["documents/doc-1/slides/1.png"])
        units = ing.ingest_pptx(_pptx_bytes(1), "doc-1", ds)
        assert units[0].classification in (
            "title_only", "text_heavy", "diagram_heavy", "hybrid"
        )

    def test_existing_pngs_are_not_re_rendered(self):
        """The upload route rasterises in a separate task; re-rendering would
        duplicate the work and the S3 writes on every extraction."""
        ds = _DocStore(present=["documents/doc-1/slides/1.png"])
        ing.ingest_pptx(_pptx_bytes(2), "doc-1", ds, converted_pdf_bytes=_pdf_bytes(2))
        assert ds.uploads == []

    def test_missing_pngs_are_rendered_from_the_converted_pdf(self):
        """The rasterise task races this trigger, so absence is expected — and
        multimodal is the whole point, so recover locally rather than degrade."""
        ds = _DocStore(present=[])
        units = ing.ingest_pptx(
            _pptx_bytes(2), "doc-1", ds, converted_pdf_bytes=_pdf_bytes(2)
        )
        assert [k for k, _ in ds.uploads] == [
            "documents/doc-1/slides/1.png",
            "documents/doc-1/slides/2.png",
        ]
        assert len(units) == 2

    def test_no_converted_pdf_degrades_to_text_only(self):
        ds = _DocStore(present=[])
        units = ing.ingest_pptx(_pptx_bytes(2), "doc-1", ds, converted_pdf_bytes=None)
        assert ds.uploads == []
        assert len(units) == 2  # still ingested, just without fresh images

    def test_a_rasterisation_failure_still_returns_units(self):
        """Best-effort: losing images is bad, losing the deck is worse."""
        ds = _DocStore(present=[], upload_error=RuntimeError("S3 down"))
        units = ing.ingest_pptx(
            _pptx_bytes(2), "doc-1", ds, converted_pdf_bytes=_pdf_bytes(2)
        )
        assert len(units) == 2

    def test_shape_counts_are_recorded(self):
        ds = _DocStore(present=["documents/doc-1/slides/1.png"])
        units = ing.ingest_pptx(_pptx_bytes(1), "doc-1", ds)
        assert units[0].n_shapes >= 1
        assert units[0].n_diagram_shapes >= 0

    def test_diagram_shape_count_never_goes_negative(self):
        ds = _DocStore(present=["documents/doc-1/slides/1.png"])
        units = ing.ingest_pptx(_pptx_bytes(1, with_table=True), "doc-1", ds)
        assert all(u.n_diagram_shapes >= 0 for u in units)


class TestRasterisePdfToSlideKeys:
    def test_uploads_one_png_per_page_one_based(self):
        ds = _DocStore()
        ing._rasterise_pdf_to_slide_keys(_pdf_bytes(3), "doc-1", ds)
        assert [k for k, _ in ds.uploads] == [
            "documents/doc-1/slides/1.png",
            "documents/doc-1/slides/2.png",
            "documents/doc-1/slides/3.png",
        ]

    def test_pngs_are_non_empty(self):
        ds = _DocStore()
        ing._rasterise_pdf_to_slide_keys(_pdf_bytes(1), "doc-1", ds)
        assert ds.uploads[0][1] > 0


class TestIngestPdf:
    def test_produces_units_from_sections(self):
        ds = _DocStore()
        units = ing.ingest_pdf(_pdf_bytes(2), "doc-1", ds)
        assert units
        assert all(u.source_format == "pdf" for u in units)
        assert all(u.classification == "pdf_section" for u in units)

    def test_page_pngs_are_uploaded_for_multimodal_extraction(self):
        """render_pages_to_png was validator-only; wiring it here is what fixed
        PDF diagram-blindness (vector diagrams emit no embedded image)."""
        ds = _DocStore()
        ing.ingest_pdf(_pdf_bytes(2), "doc-1", ds)
        keys = [k for k, _ in ds.uploads]
        assert "documents/doc-1/pages/1.png" in keys
        assert "documents/doc-1/pages/2.png" in keys

    def test_units_reference_their_page_image(self):
        ds = _DocStore()
        units = ing.ingest_pdf(_pdf_bytes(1), "doc-1", ds)
        assert any(u.image_ref == "documents/doc-1/pages/1.png" for u in units)

    def test_an_upload_failure_leaves_that_page_text_only(self):
        ds = _DocStore(upload_error=RuntimeError("S3 down"))
        units = ing.ingest_pdf(_pdf_bytes(2), "doc-1", ds)
        assert units  # still ingested
        assert all(u.image_ref is None for u in units)

    def test_text_is_carried_into_the_unit(self):
        ds = _DocStore()
        units = ing.ingest_pdf(_pdf_bytes(1, text="Multi-AZ required"), "doc-1", ds)
        assert any("Multi-AZ" in (u.text or "") for u in units)

    def test_corrupt_pdf_bytes_raise(self):
        """Unlike the image path, unreadable input is not recoverable — better to
        surface it than to record an empty extraction as success."""
        with pytest.raises(Exception):
            ing.ingest_pdf(b"not a pdf at all", "doc-1", _DocStore())


class TestIngestDispatch:
    def test_pptx_documents_route_to_the_pptx_ingester(self):
        ds = _DocStore(present=["documents/doc-1/slides/1.png"])
        units = ing.ingest(
            _doc(SourceFormat.PPTX), ds, pptx_bytes=_pptx_bytes(2)
        )
        assert all(u.source_format == "pptx" for u in units)

    def test_pdf_documents_route_to_the_pdf_ingester(self):
        ds = _DocStore()
        units = ing.ingest(_doc(SourceFormat.PDF), ds, pdf_bytes=_pdf_bytes(1))
        assert all(u.source_format == "pdf" for u in units)

    def test_pptx_without_bytes_is_rejected_loudly(self):
        with pytest.raises(ValueError, match="pptx_bytes"):
            ing.ingest(_doc(SourceFormat.PPTX), _DocStore())

    def test_pdf_without_bytes_is_rejected_loudly(self):
        with pytest.raises(ValueError, match="pdf_bytes"):
            ing.ingest(_doc(SourceFormat.PDF), _DocStore())

    def test_for_pptx_the_pdf_bytes_are_the_converted_deck(self):
        """Same kwarg, different meaning per format — passing it as the source
        document instead of the converted render would ingest the wrong thing."""
        ds = _DocStore(present=[])
        units = ing.ingest(
            _doc(SourceFormat.PPTX),
            ds,
            pptx_bytes=_pptx_bytes(2),
            pdf_bytes=_pdf_bytes(2),
        )
        # Slides were rasterised from the converted PDF, and the units are
        # still the PPTX's slides (not the PDF's pages).
        assert all(u.source_format == "pptx" for u in units)
        assert [k for k, _ in ds.uploads] == [
            "documents/doc-1/slides/1.png",
            "documents/doc-1/slides/2.png",
        ]

    def test_a_document_without_source_format_defaults_to_pdf(self):
        class _Bare:
            document_id = "doc-1"

        units = ing.ingest(_Bare(), _DocStore(), pdf_bytes=_pdf_bytes(1))
        assert all(u.source_format == "pdf" for u in units)
